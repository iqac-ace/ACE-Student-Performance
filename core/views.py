from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib import messages
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Avg, Count, Q
from openpyxl import load_workbook, Workbook

from .models import *
from .forms import *
from .utils import student_spi


# ============================================================
# USER ROLE
# ============================================================

def role(user):
    """
    Return the user's application role.

    Superuser is treated as IQAC.
    """

    if user.is_superuser:
        return "IQAC"

    try:
        return user.profile.role
    except Exception:
        return "FACULTY"


# ============================================================
# DEPARTMENT DATA SCOPE
# ============================================================

def dept_scope(user, qs):
    """
    IQAC/Superuser can see all departments.
    HOD/Faculty can see only their department.
    """

    r = role(user)

    if r == "IQAC" or user.is_superuser:
        return qs

    try:
        return qs.filter(
            department=user.profile.department
        )
    except Exception:
        return qs.none()


# ============================================================
# DASHBOARD
# ============================================================

@login_required
def dashboard(request):

    students_qs = dept_scope(
        request.user,
        Student.objects.filter(active=True)
    )

    dept_qs = Department.objects.filter(active=True)

    if role(request.user) != "IQAC" and not request.user.is_superuser:
        try:
            dept_qs = dept_qs.filter(
                id=request.user.profile.department_id
            )
        except Exception:
            dept_qs = dept_qs.none()

    cards = {
        "students": students_qs.count(),

        "departments": dept_qs.count(),

        "pending": ActivityEvidence.objects.filter(
            student__in=students_qs,
            status="PENDING"
        ).count(),

        "semester_results": SemesterResult.objects.filter(
            student__in=students_qs
        ).count(),
    }

    dept_stats = []

    for department in dept_qs:

        students_in_dept = Student.objects.filter(
            department=department,
            active=True
        )

        avg12 = students_in_dept.aggregate(
            value=Avg("class12_percentage")
        )["value"] or 0

        avgcgpa = SemesterResult.objects.filter(
            student__in=students_in_dept,
            status__in=[
                "HOD_APPROVED",
                "IQAC_APPROVED"
            ]
        ).aggregate(
            value=Avg("cgpa")
        )["value"] or 0

        dept_stats.append({
            "name": department.code,
            "students": students_in_dept.count(),
            "avg12": round(float(avg12), 1),
            "avgcgpa": round(float(avgcgpa), 2),
        })

    recent = (
        SemesterResult.objects
        .filter(student__in=students_qs)
        .select_related("student")
        .order_by("-id")[:8]
    )

    return render(
        request,
        "core/dashboard.html",
        {
            "cards": cards,
            "dept_stats": dept_stats,
            "recent": recent,
            "role": role(request.user),
        }
    )


# ============================================================
# STUDENT SEARCH / LIST
# ============================================================

@login_required
def students(request):

    q = request.GET.get("q", "").strip()

    qs = dept_scope(
        request.user,
        Student.objects
        .select_related("department")
        .filter(active=True)
    )

    if q:
        qs = qs.filter(
            Q(name__icontains=q)
            |
            Q(register_number__icontains=q)
            |
            Q(department__name__icontains=q)
        )

    return render(
        request,
        "core/students.html",
        {
            "students": qs[:200],
            "q": q,
        }
    )


# ============================================================
# STUDENT COMPLETE DETAIL
# ============================================================

@login_required
def student_detail(request, pk):

    s = get_object_or_404(
        dept_scope(
            request.user,
            Student.objects.select_related("department")
        ),
        pk=pk
    )

    marks = (
        s.subject_marks
        .select_related("subject")
        .order_by(
            "subject__semester",
            "subject__code"
        )
    )

    semesters = s.semester_results.order_by("semester")

    activities = (
        s.activities
        .select_related("parameter")
        .order_by("-activity_date")
    )

    scores, total, level = student_spi(s)

    trend = [
        {
            "semester": result.semester,
            "sgpa": float(result.sgpa or 0),
            "cgpa": float(result.cgpa or 0),
        }
        for result in semesters
    ]

    return render(
        request,
        "core/student_detail.html",
        {
            "s": s,
            "marks": marks,
            "semesters": semesters,
            "activities": activities,
            "scores": scores,
            "spi": total,
            "level": level,
            "trend": trend,
        }
    )


# ============================================================
# ADD STUDENT
# ============================================================

@login_required
def student_add(request):

    # Only IQAC and HOD can add students
    if role(request.user) not in ["IQAC", "HOD"]:
        return HttpResponseForbidden(
            "Only IQAC Coordinator or HOD can add students."
        )

    form = StudentForm(
        request.POST or None
    )

    if form.is_valid():

        student = form.save(commit=False)

        # HOD can add only to own department
        if role(request.user) == "HOD":

            try:
                student.department = (
                    request.user.profile.department
                )
            except Exception:
                return HttpResponseForbidden(
                    "HOD department is not configured."
                )

        student.save()

        # Audit log
        try:
            AuditLog.objects.create(
                actor=request.user,
                action="Student created",
                entity="Student",
                entity_id=str(student.id),
                details=(
                    f"Created student: "
                    f"{student.register_number} - "
                    f"{student.name}"
                )
            )
        except Exception:
            pass

        messages.success(
            request,
            f"Student {student.name} added successfully."
        )

        return redirect(
            "student_detail",
            pk=student.pk
        )

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Add Student",
        }
    )


# ============================================================
# EDIT STUDENT
# ============================================================

@login_required
def student_edit(request, pk):

    # Only IQAC and HOD can edit
    if role(request.user) not in ["IQAC", "HOD"]:
        return HttpResponseForbidden(
            "Only IQAC Coordinator or HOD can edit student details."
        )

    student = get_object_or_404(
        Student,
        pk=pk
    )

    # HOD can edit only own department
    if role(request.user) == "HOD":

        try:
            if (
                student.department_id
                != request.user.profile.department_id
            ):
                return HttpResponseForbidden(
                    "You cannot edit students from another department."
                )
        except Exception:
            return HttpResponseForbidden(
                "HOD department is not configured."
            )

    form = StudentForm(
        request.POST or None,
        instance=student
    )

    if form.is_valid():

        updated_student = form.save(
            commit=False
        )

        # HOD cannot transfer student
        # to another department
        if role(request.user) == "HOD":

            updated_student.department = (
                request.user.profile.department
            )

        updated_student.save()

        # Audit log
        try:
            AuditLog.objects.create(
                actor=request.user,
                action="Student details updated",
                entity="Student",
                entity_id=str(updated_student.id),
                details=(
                    f"Updated student: "
                    f"{updated_student.register_number} - "
                    f"{updated_student.name}"
                )
            )
        except Exception:
            pass

        messages.success(
            request,
            "Student details updated successfully."
        )

        return redirect(
            "student_detail",
            pk=updated_student.pk
        )

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": (
                f"Edit Student - "
                f"{student.name}"
            ),
        }
    )


# ============================================================
# DEPARTMENTS
# ============================================================

@login_required
def departments(request):

    if role(request.user) == "IQAC":
        qs = Department.objects.all()
    else:
        try:
            qs = Department.objects.filter(
                id=request.user.profile.department_id
            )
        except Exception:
            qs = Department.objects.none()

    return render(
        request,
        "core/departments.html",
        {
            "departments": qs
        }
    )


@login_required
def department_add(request):

    if role(request.user) != "IQAC":
        return HttpResponseForbidden(
            "IQAC Coordinator only."
        )

    form = DepartmentForm(
        request.POST or None
    )

    if form.is_valid():

        department = form.save()

        messages.success(
            request,
            f"Department {department.name} created successfully."
        )

        return redirect("departments")

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Create Department",
        }
    )


# ============================================================
# CREATE FACULTY / HOD LOGIN
# ============================================================

@login_required
def staff_add(request):

    if role(request.user) != "IQAC":
        return HttpResponseForbidden(
            "IQAC Coordinator only."
        )

    form = StaffCreateForm(
        request.POST or None
    )

    if form.is_valid():

        username = form.cleaned_data["username"]

        if User.objects.filter(
            username=username
        ).exists():

            form.add_error(
                "username",
                "Username already exists."
            )

        else:

            user = User.objects.create_user(
                username=username,
                email=form.cleaned_data["email"],
                password=form.cleaned_data["password"],
                first_name=form.cleaned_data["first_name"],
                last_name=form.cleaned_data["last_name"],
            )

            Profile.objects.create(
                user=user,
                role=form.cleaned_data["role"],
                department=form.cleaned_data["department"],
            )

            messages.success(
                request,
                "Staff login created successfully."
            )

            return redirect("dashboard")

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Create Faculty / HOD Login",
        }
    )


# ============================================================
# SUBJECTS
# ============================================================

@login_required
def subjects(request):

    qs = (
        Subject.objects
        .select_related("department")
        .order_by(
            "department",
            "semester",
            "code"
        )
    )

    if role(request.user) != "IQAC":

        try:
            qs = qs.filter(
                department=request.user.profile.department
            )
        except Exception:
            qs = qs.none()

    return render(
        request,
        "core/subjects.html",
        {
            "subjects": qs
        }
    )


@login_required
def subject_add(request):

    if role(request.user) not in ["IQAC", "HOD"]:
        return HttpResponseForbidden(
            "Not authorized."
        )

    form = SubjectForm(
        request.POST or None
    )

    if form.is_valid():

        subject = form.save()

        messages.success(
            request,
            f"Subject {subject.code} added successfully."
        )

        return redirect("subjects")

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Add Subject",
        }
    )


# ============================================================
# EXCEL MARKS TEMPLATE
# ============================================================

@login_required
def marks_template(request):

    wb = Workbook()

    ws = wb.active
    ws.title = "Marks"

    ws.append([
        "Register Number",
        "Subject Code",
        "Internal",
        "External"
    ])

    ws.append([
        "21CSE001",
        "CS401",
        "35",
        "54"
    ])

    response = HttpResponse(
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="ACE_Marks_Upload_Template.xlsx"'
    )

    wb.save(response)

    return response


# ============================================================
# EXCEL MARKS UPLOAD
# ============================================================

@login_required
def marks_upload(request):

    form = ExcelUploadForm(
        request.POST or None,
        request.FILES or None
    )

    if request.method == "POST" and form.is_valid():

        wb = load_workbook(
            form.cleaned_data["file"],
            data_only=True
        )

        ws = wb.active

        success_count = 0
        errors = []

        for rowno, row in enumerate(
            ws.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            if not row or not row[0]:
                continue

            try:

                reg = row[0]
                subject_code = row[1]
                internal = row[2]
                external = row[3]

                student = Student.objects.get(
                    register_number=str(
                        reg
                    ).strip()
                )

                # Faculty/HOD restricted to department
                if role(request.user) != "IQAC":

                    try:
                        if (
                            student.department_id
                            != request.user.profile.department_id
                        ):
                            raise ValueError(
                                "Student is outside your department."
                            )
                    except Exception:
                        raise ValueError(
                            "User department is not configured."
                        )

                subject = Subject.objects.get(
                    department=student.department,
                    code=str(subject_code).strip()
                )

                SubjectMark.objects.update_or_create(

                    student=student,

                    subject=subject,

                    defaults={
                        "internal": internal or 0,
                        "external": external or 0,
                        "status": "PENDING",
                        "uploaded_by": request.user,
                    }
                )

                success_count += 1

            except Exception as e:

                errors.append(
                    f"Row {rowno}: {e}"
                )

        messages.success(
            request,
            f"{success_count} mark rows imported."
        )

        for error in errors[:8]:

            messages.warning(
                request,
                error
            )

        return redirect(
            "marks_upload"
        )

    return render(
        request,
        "core/upload.html",
        {
            "form": form
        }
    )


# ============================================================
# SEMESTER RESULT
# ============================================================

@login_required
def semester_add(request):

    form = SemesterResultForm(
        request.POST or None
    )

    if form.is_valid():

        obj = form.save(
            commit=False
        )

        # Department restriction
        if role(request.user) != "IQAC":

            try:
                if (
                    obj.student.department_id
                    != request.user.profile.department_id
                ):
                    return HttpResponseForbidden(
                        "Wrong department."
                    )
            except Exception:
                return HttpResponseForbidden(
                    "User department is not configured."
                )

        obj.uploaded_by = request.user
        obj.status = "PENDING"

        obj.save()

        messages.success(
            request,
            "Semester result saved for approval."
        )

        return redirect(
            "student_detail",
            pk=obj.student_id
        )

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Add Semester Result",
        }
    )


# ============================================================
# ACTIVITY / CULTURAL / TECHNICAL RECORD
# ============================================================

@login_required
def activity_add(request):

    form = ActivityEvidenceForm(
        request.POST or None,
        request.FILES or None
    )

    if form.is_valid():

        obj = form.save(
            commit=False
        )

        # Department restriction
        if role(request.user) != "IQAC":

            try:
                if (
                    obj.student.department_id
                    != request.user.profile.department_id
                ):
                    return HttpResponseForbidden(
                        "Wrong department."
                    )
            except Exception:
                return HttpResponseForbidden(
                    "User department is not configured."
                )

        obj.created_by = request.user
        obj.status = "PENDING"

        ratios = {
            0: 0,
            1: 0.37,
            2: 0.62,
            3: 0.82,
            4: 0.95,
        }

        obj.points = round(
            obj.parameter.max_points
            * ratios.get(obj.level, 0),
            2
        )

        obj.save()

        messages.success(
            request,
            "Activity/evidence submitted."
        )

        return redirect(
            "student_detail",
            pk=obj.student_id
        )

    return render(
        request,
        "core/form.html",
        {
            "form": form,
            "title": "Add Activity / Evidence",
        }
    )


# ============================================================
# APPROVALS
# ============================================================

@login_required
def approvals(request):

    r = role(request.user)

    if r not in ["IQAC", "HOD"]:
        return HttpResponseForbidden(
            "Not authorized."
        )

    students_qs = dept_scope(
        request.user,
        Student.objects.all()
    )

    activities = ActivityEvidence.objects.filter(
        student__in=students_qs
    )

    semesters = SemesterResult.objects.filter(
        student__in=students_qs
    )

    marks = SubjectMark.objects.filter(
        student__in=students_qs
    )

    if r == "HOD":

        activities = activities.filter(
            status="PENDING"
        )

        semesters = semesters.filter(
            status="PENDING"
        )

        marks = marks.filter(
            status="PENDING"
        )

    else:

        activities = activities.filter(
            status="HOD_APPROVED"
        )

        semesters = semesters.filter(
            status="HOD_APPROVED"
        )

        marks = marks.filter(
            status="HOD_APPROVED"
        )

    return render(
        request,
        "core/approvals.html",
        {
            "activities": activities.select_related(
                "student",
                "parameter"
            ),

            "semesters": semesters.select_related(
                "student"
            ),

            "marks": marks.select_related(
                "student",
                "subject"
            ),

            "role": r,
        }
    )


# ============================================================
# APPROVE / REJECT
# ============================================================

@login_required
def approve(request, model, pk, decision):

    r = role(request.user)

    modelmap = {
        "activity": ActivityEvidence,
        "semester": SemesterResult,
        "mark": SubjectMark,
    }

    if model not in modelmap:

        return HttpResponseForbidden(
            "Invalid approval type."
        )

    obj = get_object_or_404(
        modelmap[model],
        pk=pk
    )

    # HOD department security
    if r == "HOD":

        try:

            if (
                obj.student.department_id
                != request.user.profile.department_id
            ):
                return HttpResponseForbidden(
                    "You cannot approve records from another department."
                )

        except Exception:

            return HttpResponseForbidden(
                "HOD department is not configured."
            )

    # HOD approval
    if (
        r == "HOD"
        and obj.status == "PENDING"
    ):

        if decision == "approve":
            obj.status = "HOD_APPROVED"
        else:
            obj.status = "REJECTED"

    # IQAC approval
    elif (
        r == "IQAC"
        and obj.status == "HOD_APPROVED"
    ):

        if decision == "approve":
            obj.status = "IQAC_APPROVED"
        else:
            obj.status = "REJECTED"

    else:

        return HttpResponseForbidden(
            "Invalid approval stage."
        )

    obj.save()

    try:

        AuditLog.objects.create(
            actor=request.user,
            action=decision,
            entity=model,
            entity_id=str(pk),
            details=obj.status
        )

    except Exception:
        pass

    messages.success(
        request,
        f"Record {decision}d successfully."
    )

    return redirect(
        "approvals"
    )


# ============================================================
# REPORTS
# ============================================================

@login_required
def reports(request):

    departments = Department.objects.filter(
        active=True
    )

    rows = []

    if role(request.user) != "IQAC":

        try:
            departments = departments.filter(
                id=request.user.profile.department_id
            )
        except Exception:
            departments = departments.none()

    for department in departments:

        students_in_dept = Student.objects.filter(
            department=department,
            active=True
        )

        avg12 = students_in_dept.aggregate(
            value=Avg("class12_percentage")
        )["value"] or 0

        avgcgpa = SemesterResult.objects.filter(
            student__in=students_in_dept,
            status="IQAC_APPROVED"
        ).aggregate(
            value=Avg("cgpa")
        )["value"] or 0

        rows.append({
            "department": department,
            "students": students_in_dept.count(),
            "avg12": avg12,
            "avgcgpa": avgcgpa,
        })

    return render(
        request,
        "core/reports.html",
        {
            "rows": rows
        }
    )